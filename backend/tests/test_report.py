import io
import tempfile
import unittest
from pathlib import Path
from docx import Document
from scripts.create_template import create_template
from api.report_generator import generate_inspection_report, generate_basic_report, generate_word_report
class ReportTests(unittest.TestCase):
    def test_synthetic_report(self):
        create_template()
        result=generate_inspection_report({'inspector_name':'Demo Reviewer','contractor_name':'Example Supplier','site_address':'Synthetic site','notes':'A synthetic observation.','defects_with_photos':[], 'general_photos':[]})
        if hasattr(result,'getvalue'):result=result.getvalue()
        doc=Document(io.BytesIO(result))
        text='\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('Demo Reviewer',text)
        self.assertIn('Synthetic site',text)
        self.assertNotIn('{{ Inspector_Name',text)
    def test_clean_template(self):
        with tempfile.TemporaryDirectory() as directory:
            path=create_template(Path(directory)/'template.docx')
            doc=Document(path)
            self.assertEqual(doc.core_properties.author,'')
            self.assertIn('Field observation report','\n'.join(p.text for p in doc.paragraphs))
    def test_fallback_export_uses_generic_report_labels(self):
        result = generate_basic_report({'inspector_name': 'Demo Reviewer', 'notes': 'Synthetic note'})
        doc = Document(io.BytesIO(result))
        text = '\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('Field Observation Report', text)
        self.assertIn('Reviewer Details', text)
        self.assertIn('Demo Reviewer', text)

    def test_analysis_export_uses_generic_title(self):
        result = generate_word_report({'overall_assessment': 'Synthetic observation for review.'})
        doc = Document(io.BytesIO(result))
        text = '\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('Field Notes Analysis Report', text)
        self.assertIn('Synthetic observation for review.', text)
if __name__=='__main__':unittest.main()
