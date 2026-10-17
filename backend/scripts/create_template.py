"""Generate a clean, generic report template; no original report binary is included."""
from pathlib import Path
from docx import Document

def create_template(path=None):
    path = Path(path) if path else Path(__file__).resolve().parent.parent / 'templates/template.docx'
    path.parent.mkdir(parents=True, exist_ok=True)
    doc=Document()
    doc.core_properties.author=''
    doc.core_properties.last_modified_by=''
    doc.add_heading('Field observation report', 0)
    doc.add_paragraph('AI-assisted draft. All observations require human review. Demonstration categories are not standards.')
    for label, key in [('Reviewer','Inspector_Name'),('Supplier','Contractor_Name'),('Location','Site_Address'),('Date','Inspection_Date')]:
        doc.add_paragraph(label + ': {{ ' + key + ' }}')
    doc.add_heading('Main photo', 1)
    doc.add_paragraph('{{ Main_Site_Image }}')
    doc.add_heading('Observations',1)
    doc.add_paragraph('{{ Notes }}')
    doc.save(path)
    return path

if __name__=='__main__':
    print(create_template())
